"""Utility to merge exercise progression data into the SQLite database with diff exports."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import shutil
import sqlite3
import sys
import time
import traceback
from dataclasses import dataclass, field as dc_field
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

from muscle_norm import (
    DEFAULT_ACRONYMS as MUSCLE_DEFAULT_ACRONYMS,
    get_alias_acronyms as get_muscle_alias_acronyms,
    load_alias_map as load_muscle_alias_map,
    normalize_list as normalize_muscle_list,
    normalize_single as normalize_muscle_single,
    split_list as split_muscle_list,
)

try:
    import pandas as pd  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    pd = None  # type: ignore

DEFAULT_CSV_PATH = os.path.join("data", "merge_preview.csv")
DEFAULT_DB_PATH = os.path.join("data", "database.db")
DEFAULT_LOG_PATH = os.path.join("logs", "merge_progress.md")

PRIMARY_KEY = "exercise_name"
TARGET_COLUMNS: Sequence[str] = (
    PRIMARY_KEY,
    "primary_muscle_group",
    "secondary_muscle_group",
    "tertiary_muscle_group",
    "advanced_isolated_muscles",
    "utility",
    "grips",
    "stabilizers",
    "synergists",
    "force",
    "equipment",
    "mechanic",
    "difficulty",
)
LIST_COLUMNS: Set[str] = {
    "advanced_isolated_muscles",
    "grips",
    "stabilizers",
    "synergists",
}
SCALAR_COLUMNS: List[str] = [column for column in TARGET_COLUMNS if column not in LIST_COLUMNS and column != PRIMARY_KEY]
DATA_COLUMNS: Sequence[str] = tuple(column for column in TARGET_COLUMNS if column != PRIMARY_KEY)

MUSCLE_SCALAR_COLUMNS: Sequence[str] = (
    "primary_muscle_group",
    "secondary_muscle_group",
    "tertiary_muscle_group",
)
MUSCLE_LIST_COLUMNS: Sequence[str] = ("advanced_isolated_muscles",)
MUSCLE_COLUMNS: Set[str] = set(MUSCLE_SCALAR_COLUMNS) | set(MUSCLE_LIST_COLUMNS)

LIST_EMPTY_TOKENS = {"", "none", "n/a", "na", "no significant stabilizers"}

ALLOWED_SCALAR_POLICIES: Set[str] = {"db", "csv", "prefer_longer", "prefer_shorter", "warn"}
ALLOWED_LIST_POLICIES: Set[str] = {"db", "csv", "union", "replace"}

CsvRow = Dict[str, Optional[str]]


@dataclass
class ColumnDiff:
    field: str
    old_value: Optional[str]
    new_value: Optional[str]
    action: str
    tokens_added: List[str] = dc_field(default_factory=list)


@dataclass
class RowDiff:
    exercise_name: str
    change_type: str
    columns: List[ColumnDiff] = dc_field(default_factory=list)

    @property
    def columns_changed(self) -> List[str]:
        return [column.field for column in self.columns]


@dataclass
class MergeSummary:
    timestamp: dt.datetime
    csv_path: str
    db_path: str
    log_path: str
    dryrun: bool
    fix_names: bool
    scalar_policy: str
    scalar_overrides: Dict[str, str]
    list_policy: str
    list_overrides: Dict[str, str]
    diff_top_k: int
    backup_path: Optional[str] = None
    diff_output_path: Optional[str] = None
    diff_format: Optional[str] = None
    name_trim_status: str = "not requested"
    name_trim_csv: Optional[str] = None
    csv_rows: int = 0
    consolidated_records: int = 0
    baseline_rows: int = 0
    inserted: int = 0
    updated: int = 0
    no_ops: int = 0
    conflicts: int = 0
    duration_seconds: float = 0.0
    diff_entries: List[RowDiff] = dc_field(default_factory=list)
    top_updates: List[RowDiff] = dc_field(default_factory=list)
    skipped_rows: List[str] = dc_field(default_factory=list)
    errors: Optional[str] = None
    muscle_alias_file: Optional[str] = None
    muscle_acronyms: Set[str] = dc_field(default_factory=set)


@dataclass
class PolicyConfig:
    scalar_default: str
    scalar_overrides: Dict[str, str]
    list_default: str
    list_overrides: Dict[str, str]

    def scalar_for(self, field: str) -> str:
        return self.scalar_overrides.get(field, self.scalar_default)

    def list_for(self, field: str) -> str:
        return self.list_overrides.get(field, self.list_default)


def normalize_header(header: str) -> str:
    return "".join(ch for ch in header.lower() if ch.isalnum())


def normalize_scalar(value: Optional[object]) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def normalize_list_tokens(value: Optional[object]) -> List[str]:
    if value is None:
        return []
    tokens: List[str] = []
    for chunk in str(value).replace(",", ";").split(";"):
        token = chunk.strip()
        if not token:
            continue
        lowered = token.lower()
        if lowered in LIST_EMPTY_TOKENS:
            continue
        tokens.append(token)
    return tokens


def canonicalize_list_tokens(tokens: Iterable[str]) -> List[str]:
    unique: List[str] = []
    seen: Set[str] = set()
    for token in tokens:
        lowered = token.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        unique.append(token)
    unique.sort(key=lambda item: item.lower())
    return unique


def join_list_tokens(tokens: Iterable[str]) -> Optional[str]:
    canonical = canonicalize_list_tokens(tokens)
    return "; ".join(canonical) if canonical else None


def parse_list_field(value: Optional[object]) -> Optional[str]:
    return join_list_tokens(normalize_list_tokens(value))


def merge_list_strings(existing: Optional[str], incoming: Optional[str]) -> Optional[str]:
    tokens: List[str] = []
    tokens.extend(normalize_list_tokens(existing))
    tokens.extend(normalize_list_tokens(incoming))
    return join_list_tokens(tokens)


def values_equal(original: Optional[str], updated: Optional[str]) -> bool:
    return original == updated


def apply_muscle_normalization(record: CsvRow, alias_map: Dict[str, str], acronyms: Set[str]) -> None:
    for column in MUSCLE_SCALAR_COLUMNS:
        if column in record:
            record[column] = normalize_muscle_single(record.get(column), alias_map, acronyms)
    for column in MUSCLE_LIST_COLUMNS:
        if column in record:
            record[column] = normalize_muscle_list(record.get(column), alias_map, acronyms)


def compute_tokens_added(
    db_value: Optional[str],
    new_value: Optional[str],
    alias_map: Dict[str, str],
    acronyms: Set[str],
) -> List[str]:
    if not new_value:
        return []
    final_tokens = split_muscle_list(new_value)
    db_normalized = normalize_muscle_list(db_value, alias_map, acronyms)
    existing = {token.lower() for token in split_muscle_list(db_normalized)} if db_normalized else set()
    return [token for token in final_tokens if token.lower() not in existing]


def build_header_map(headers: Sequence[str]) -> Dict[str, str]:
    normalized_targets = {normalize_header(column): column for column in TARGET_COLUMNS}
    header_map: Dict[str, str] = {}
    for header in headers:
        key = normalize_header(header)
        target = normalized_targets.get(key)
        if target and target not in header_map:
            header_map[target] = header
    return header_map


def load_csv_records(csv_path: str) -> Tuple[Dict[str, CsvRow], List[str], int]:
    if not os.path.exists(csv_path):
        raise FileNotFoundError(f"CSV file not found: {csv_path}")

    records: Dict[str, CsvRow] = {}
    skipped: List[str] = []
    total_rows = 0

    def process(row_data: Dict[str, object], row_number: int, header_map: Dict[str, str]) -> None:
        nonlocal total_rows
        total_rows += 1
        name_header = header_map.get(PRIMARY_KEY, PRIMARY_KEY)
        raw_name_value = row_data.get(name_header)
        raw_name = "" if raw_name_value is None else str(raw_name_value)
        trimmed_name = raw_name.strip()
        if not trimmed_name:
            skipped.append(f"Row {row_number} (missing exercise_name)")
            return

        record: CsvRow = {PRIMARY_KEY: trimmed_name}
        for column in SCALAR_COLUMNS:
            source = header_map.get(column, column)
            record[column] = normalize_scalar(row_data.get(source))
        for column in LIST_COLUMNS:
            source = header_map.get(column, column)
            record[column] = parse_list_field(row_data.get(source))

        existing = records.get(trimmed_name)
        if existing is None:
            records[trimmed_name] = record
        else:
            for column in SCALAR_COLUMNS:
                if existing.get(column) is None and record.get(column) is not None:
                    existing[column] = record[column]
            for column in LIST_COLUMNS:
                existing[column] = merge_list_strings(existing.get(column), record.get(column))

    if pd is not None:  # pragma: no cover - optional dependency path
        frame = pd.read_csv(csv_path, dtype=str, keep_default_na=False, na_values=[])
        header_map = build_header_map(list(frame.columns))
        for index, row in enumerate(frame.to_dict(orient="records"), start=2):
            cast_row: Dict[str, object] = {str(key): value for key, value in row.items()}
            process(cast_row, index, header_map)
    else:
        with open(csv_path, "r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames is None:
                raise ValueError("CSV file is missing a header row")
            header_map = build_header_map(reader.fieldnames)
            for index, row in enumerate(reader, start=2):
                cast_row = {str(key): value for key, value in row.items()}
                process(cast_row, index, header_map)

    return records, skipped, total_rows


def ensure_exercises_table(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS exercises (
            exercise_name TEXT PRIMARY KEY,
            primary_muscle_group TEXT,
            secondary_muscle_group TEXT,
            tertiary_muscle_group TEXT,
            advanced_isolated_muscles TEXT,
            utility TEXT,
            grips TEXT,
            stabilizers TEXT,
            synergists TEXT,
            force TEXT,
            equipment TEXT,
            mechanic TEXT,
            difficulty TEXT
        )
        """
    )


def backup_database(db_path: str) -> str:
    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = f"{db_path}.backup_{timestamp}"
    shutil.copy2(db_path, backup_path)
    return backup_path


def fetch_existing_row(conn: sqlite3.Connection, name: str) -> Optional[sqlite3.Row]:
    cursor = conn.execute(
        f"SELECT {', '.join(TARGET_COLUMNS)} FROM exercises WHERE {PRIMARY_KEY} = ?",
        (name,),
    )
    return cursor.fetchone()


def insert_row(conn: sqlite3.Connection, record: CsvRow) -> None:
    columns = list(TARGET_COLUMNS)
    placeholders = ", ".join("?" for _ in columns)
    values = [record.get(column) for column in columns]
    conn.execute(
        f"INSERT INTO exercises ({', '.join(columns)}) VALUES ({placeholders})",
        values,
    )


def update_row(conn: sqlite3.Connection, record: CsvRow) -> None:
    columns = [column for column in TARGET_COLUMNS if column != PRIMARY_KEY]
    assignments = ", ".join(f"{column} = ?" for column in columns)
    params: List[Optional[str]] = [record.get(column) for column in columns]
    params.append(record[PRIMARY_KEY])
    conn.execute(
        f"UPDATE exercises SET {assignments} WHERE {PRIMARY_KEY} = ?",
        params,
    )


def build_insert_record(
    csv_record: CsvRow,
    alias_map: Dict[str, str],
    acronyms: Set[str],
) -> Tuple[CsvRow, RowDiff]:
    name = csv_record.get(PRIMARY_KEY)
    if not name:
        raise ValueError("CSV record missing exercise_name")

    record: CsvRow = {PRIMARY_KEY: name}
    row_diff = RowDiff(exercise_name=name, change_type="Inserted")

    for column in DATA_COLUMNS:
        raw_value = csv_record.get(column)
        if column in MUSCLE_SCALAR_COLUMNS:
            value = normalize_muscle_single(raw_value, alias_map, acronyms)
        elif column in MUSCLE_LIST_COLUMNS:
            value = normalize_muscle_list(raw_value, alias_map, acronyms)
        else:
            value = raw_value
        record[column] = value
        if value is None:
            continue
        if column in LIST_COLUMNS:
            tokens = split_muscle_list(value) if column in MUSCLE_LIST_COLUMNS else canonicalize_list_tokens(normalize_list_tokens(value))
            row_diff.columns.append(
                ColumnDiff(
                    field=column,
                    old_value=None,
                    new_value=value,
                    action="used CSV (list)",
                    tokens_added=tokens,
                )
            )
        else:
            row_diff.columns.append(
                ColumnDiff(
                    field=column,
                    old_value=None,
                    new_value=value,
                    action="used CSV",
                )
            )

    return record, row_diff


def resolve_scalar(
    db_value: Optional[str],
    csv_value: Optional[str],
    policy: str,
) -> Tuple[Optional[str], str, bool, bool]:
    if policy not in ALLOWED_SCALAR_POLICIES:
        raise ValueError(f"Unsupported scalar policy: {policy}")

    db_norm = normalize_scalar(db_value)
    csv_norm = normalize_scalar(csv_value)

    if csv_norm is None:
        return db_norm, "no-op", False, False
    if db_norm is None:
        return csv_norm, "used CSV", True, False
    if db_norm == csv_norm:
        return db_norm, "no-op", False, False

    if policy == "csv":
        return csv_norm, "used CSV", True, True
    if policy == "db":
        return db_norm, "kept DB", False, True
    if policy == "prefer_longer":
        if len(csv_norm) > len(db_norm):
            return csv_norm, "used CSV (longer)", True, True
        return db_norm, "kept DB (longer)", False, False
    if policy == "prefer_shorter":
        if len(csv_norm) < len(db_norm):
            return csv_norm, "used CSV (shorter)", True, True
        return db_norm, "kept DB (shorter)", False, False
    if policy == "warn":
        return db_norm, "kept DB (warn)", False, True

    raise ValueError(f"Unhandled scalar policy: {policy}")


def resolve_list(
    db_value: Optional[str],
    csv_value: Optional[str],
    policy: str,
) -> Tuple[Optional[str], List[str], str, bool]:
    if policy not in ALLOWED_LIST_POLICIES:
        raise ValueError(f"Unsupported list policy: {policy}")

    db_tokens = canonicalize_list_tokens(normalize_list_tokens(db_value))
    csv_tokens = canonicalize_list_tokens(normalize_list_tokens(csv_value))
    db_string = "; ".join(db_tokens) if db_tokens else None
    csv_string = "; ".join(csv_tokens) if csv_tokens else None

    if db_tokens == csv_tokens:
        return db_string, [], "no-op", False

    if policy == "db":
        return db_string, [], "kept DB", False

    if policy in {"csv", "replace"}:
        if csv_string is None and policy == "csv":
            return db_string, [], "no-op", False
        return csv_string, csv_tokens, "used CSV", csv_string != db_string

    if policy == "union":
        merged_tokens = canonicalize_list_tokens([*db_tokens, *csv_tokens])
        merged_string = "; ".join(merged_tokens) if merged_tokens else None
        existing_set = {token.lower() for token in db_tokens}
        tokens_added = [token for token in merged_tokens if token.lower() not in existing_set]
        return merged_string, tokens_added, "union", merged_string != db_string

    raise ValueError(f"Unhandled list policy: {policy}")


def merge_existing_record(
    existing_row: sqlite3.Row,
    csv_record: CsvRow,
    policy: PolicyConfig,
    summary: MergeSummary,
    alias_map: Dict[str, str],
    acronyms: Set[str],
) -> Tuple[CsvRow, RowDiff, bool]:
    existing_name = existing_row[PRIMARY_KEY]
    result: CsvRow = {PRIMARY_KEY: existing_name}
    row_diff = RowDiff(exercise_name=existing_name, change_type="No-op")
    changed = False

    for column in SCALAR_COLUMNS:
        db_val = existing_row[column]
        csv_val = csv_record.get(column)
        if column in MUSCLE_SCALAR_COLUMNS:
            csv_val = normalize_muscle_single(csv_val, alias_map, acronyms)
        resolved, action, _field_changed, conflict = resolve_scalar(db_val, csv_val, policy.scalar_for(column))
        final_value = (
            normalize_muscle_single(resolved, alias_map, acronyms)
            if column in MUSCLE_SCALAR_COLUMNS
            else resolved
        )
        result[column] = final_value
        field_changed = not values_equal(db_val, final_value)
        action_label = action if action != "no-op" else ("normalized" if field_changed else "no-op")
        if action_label != "no-op" or field_changed:
            row_diff.columns.append(
                ColumnDiff(
                    field=column,
                    old_value=db_val if db_val is not None else None,
                    new_value=final_value,
                    action=action_label,
                )
            )
        if field_changed:
            changed = True
        if conflict:
            summary.conflicts += 1

    for column in LIST_COLUMNS:
        db_val = existing_row[column]
        csv_val = csv_record.get(column)
        if column in MUSCLE_LIST_COLUMNS:
            csv_val = normalize_muscle_list(csv_val, alias_map, acronyms)
        resolved, tokens_added_raw, action, _field_changed = resolve_list(db_val, csv_val, policy.list_for(column))
        final_value = (
            normalize_muscle_list(resolved, alias_map, acronyms)
            if column in MUSCLE_LIST_COLUMNS
            else resolved
        )
        result[column] = final_value
        field_changed = not values_equal(db_val, final_value)
        action_label = action if action != "no-op" else ("normalized" if field_changed else "no-op")
        if column in MUSCLE_LIST_COLUMNS:
            tokens_added = compute_tokens_added(db_val, final_value, alias_map, acronyms)
        else:
            tokens_added = tokens_added_raw
        if action_label != "no-op" or field_changed or tokens_added:
            row_diff.columns.append(
                ColumnDiff(
                    field=column,
                    old_value=db_val if db_val is not None else None,
                    new_value=final_value,
                    action=action_label,
                    tokens_added=tokens_added,
                )
            )
        if field_changed:
            changed = True

    row_diff.change_type = "Updated" if changed else "No-op"
    return result, row_diff, changed


def row_diff_to_wide(row: RowDiff) -> Dict[str, str]:
    data: Dict[str, str] = {
        "change_type": row.change_type,
        "exercise_name": row.exercise_name,
        "columns_changed": ", ".join(row.columns_changed),
    }
    for column in DATA_COLUMNS:
        data[f"{column}_old"] = ""
        data[f"{column}_new"] = ""
        data[f"{column}_action"] = "no-op"
        if column in LIST_COLUMNS:
            data[f"{column}_tokens_added"] = ""

    for column_diff in row.columns:
        column_name = column_diff.field
        data[f"{column_name}_old"] = column_diff.old_value or ""
        data[f"{column_name}_new"] = column_diff.new_value or ""
        data[f"{column_name}_action"] = column_diff.action
        if column_name in LIST_COLUMNS:
            data[f"{column_name}_tokens_added"] = (
                json.dumps(column_diff.tokens_added) if column_diff.tokens_added else ""
            )

    return data


def export_diff_csv(path: str, rows: Sequence[RowDiff]) -> None:
    headers: List[str] = ["change_type", "exercise_name", "columns_changed"]
    for column in DATA_COLUMNS:
        headers.append(f"{column}_old")
        headers.append(f"{column}_new")
        headers.append(f"{column}_action")
        if column in LIST_COLUMNS:
            headers.append(f"{column}_tokens_added")

    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            wide = row_diff_to_wide(row)
            writer.writerow([wide.get(header, "") for header in headers])


def truncate(value: Optional[str], limit: int = 80) -> str:
    if not value:
        return ""
    if len(value) <= limit:
        return value
    return value[: limit - 3] + "..."


def export_diff_md(path: str, summary: MergeSummary) -> None:
    lines: List[str] = []
    lines.append(f"# Merge Diff Report ({summary.timestamp.strftime('%Y-%m-%d %H:%M:%S')})")
    lines.append("")
    lines.append(f"- Inserted: {summary.inserted}")
    lines.append(f"- Updated: {summary.updated}")
    lines.append(f"- No-ops: {summary.no_ops}")
    lines.append(f"- Conflicts: {summary.conflicts}")
    lines.append("")
    lines.append("| Change | Exercise | Field | Action | Old | New | Tokens Added |")
    lines.append("| --- | --- | --- | --- | --- | --- | --- |")
    for row in summary.diff_entries:
        if row.columns:
            for column in row.columns:
                tokens = ", ".join(column.tokens_added)
                lines.append(
                    f"| {row.change_type} | {row.exercise_name} | {column.field} | {column.action} | "
                    f"{truncate(column.old_value)} | {truncate(column.new_value)} | {truncate(tokens)} |"
                )
        else:
            lines.append(f"| {row.change_type} | {row.exercise_name} | - | no-op |  |  |  |")
    with open(path, "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines))


def export_diff_json(path: str, rows: Sequence[RowDiff]) -> None:
    payload = []
    for row in rows:
        payload.append(
            {
                "exercise_name": row.exercise_name,
                "change_type": row.change_type,
                "columns": [
                    {
                        "field": column.field,
                        "old_value": column.old_value,
                        "new_value": column.new_value,
                        "action": column.action,
                        "tokens_added": column.tokens_added,
                    }
                    for column in row.columns
                ],
            }
        )
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2)


def export_diff_xlsx(path: str, rows: Sequence[RowDiff]) -> None:
    try:
        from openpyxl import Workbook  # type: ignore
    except Exception as exc:  # pragma: no cover - optional dependency
        raise ImportError("openpyxl is required for XLSX export") from exc

    headers: List[str] = ["change_type", "exercise_name", "columns_changed"]
    for column in DATA_COLUMNS:
        headers.append(f"{column}_old")
        headers.append(f"{column}_new")
        headers.append(f"{column}_action")
        if column in LIST_COLUMNS:
            headers.append(f"{column}_tokens_added")

    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:  # pragma: no cover - safety net for type checkers
        raise RuntimeError("openpyxl failed to provide an active worksheet")
    worksheet.title = "diff"
    worksheet.append(headers)
    for row in rows:
        wide = row_diff_to_wide(row)
        worksheet.append([wide.get(header, "") for header in headers])
    workbook.save(path)


def export_diff(summary: MergeSummary, path: str, fmt: str) -> None:
    fmt_lower = fmt.lower()
    output_dir = os.path.dirname(path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    try:
        if fmt_lower == "csv":
            export_diff_csv(path, summary.diff_entries)
            summary.diff_output_path = os.path.abspath(path)
            summary.diff_format = "csv"
        elif fmt_lower == "json":
            export_diff_json(path, summary.diff_entries)
            summary.diff_output_path = os.path.abspath(path)
            summary.diff_format = "json"
        elif fmt_lower == "md":
            export_diff_md(path, summary)
            summary.diff_output_path = os.path.abspath(path)
            summary.diff_format = "md"
        elif fmt_lower == "xlsx":
            try:
                export_diff_xlsx(path, summary.diff_entries)
                summary.diff_output_path = os.path.abspath(path)
                summary.diff_format = "xlsx"
            except ImportError:
                fallback = os.path.splitext(path)[0] + ".csv"
                export_diff_csv(fallback, summary.diff_entries)
                summary.diff_output_path = os.path.abspath(fallback)
                summary.diff_format = "csv"
                print(f"Warning: openpyxl not available; diff exported as CSV to {fallback}")
        else:
            raise ValueError(f"Unsupported diff format: {fmt}")
    except Exception:
        summary.errors = summary.errors or traceback.format_exc()
        raise


def summarize_row(row: RowDiff) -> str:
    if not row.columns:
        return "no column changes"
    pieces: List[str] = []
    for column in row.columns:
        token_hint = f" ({', '.join(column.tokens_added)})" if column.tokens_added else ""
        pieces.append(f"{column.field}: {column.action}{token_hint}")
    return "; ".join(pieces)


def write_name_trim_csv(log_path: str, rows: Sequence[Tuple[str, str, str]]) -> str:
    log_dir = os.path.dirname(log_path) or "."
    os.makedirs(log_dir, exist_ok=True)
    csv_path = os.path.join(log_dir, "name_trim_candidates.csv")
    with open(csv_path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["original_name", "trimmed_name", "action"])
        for original, trimmed, action in rows:
            writer.writerow([original, trimmed, action])
    return csv_path


def perform_name_trim(db_path: str, log_path: str, dryrun: bool) -> Tuple[str, str]:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        candidates = conn.execute(
            """
            SELECT rowid, exercise_name, trim(exercise_name) AS trimmed
            FROM exercises
            WHERE exercise_name LIKE ' %' OR exercise_name LIKE '% '
            """
        ).fetchall()

        if not candidates:
            csv_path = write_name_trim_csv(log_path, [])
            return "no trim candidates", csv_path

        annotated: List[Tuple[str, str, bool]] = []
        collisions_found = False
        for row in candidates:
            original = row["exercise_name"]
            trimmed = row["trimmed"]
            has_collision = conn.execute(
                "SELECT 1 FROM exercises WHERE exercise_name = ? AND rowid <> ?",
                (trimmed, row["rowid"]),
            ).fetchone() is not None
            annotated.append((original, trimmed, has_collision))
            if has_collision:
                collisions_found = True

        if collisions_found:
            rows = [
                (original, trimmed, "collision" if has_collision else "skipped")
                for original, trimmed, has_collision in annotated
            ]
            csv_path = write_name_trim_csv(log_path, rows)
            conn.rollback()
            return "name-trim fix skipped due to potential collisions", csv_path

        if dryrun:
            rows = [(original, trimmed, "skipped") for original, trimmed, _ in annotated]
            csv_path = write_name_trim_csv(log_path, rows)
            conn.rollback()
            return "name-trim fix skipped (dryrun)", csv_path

        update_result = conn.execute(
            """
            UPDATE exercises
               SET exercise_name = trim(exercise_name)
             WHERE exercise_name LIKE ' %' OR exercise_name LIKE '% '
            """
        )
        conn.commit()
        rows = [(original, trimmed, "safe") for original, trimmed, _ in annotated]
        csv_path = write_name_trim_csv(log_path, rows)
        return f"name-trim fix applied ({update_result.rowcount} rows)", csv_path
    finally:
        conn.close()


def append_log(summary: MergeSummary) -> None:
    log_dir = os.path.dirname(summary.log_path) or "."
    os.makedirs(log_dir, exist_ok=True)

    lines: List[str] = []
    lines.append(f"## Run: {summary.timestamp.strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("- Paths:")
    lines.append(f"  - CSV: {summary.csv_path}")
    lines.append(f"  - DB: {summary.db_path}")
    if summary.backup_path:
        lines.append(f"  - Backup: {summary.backup_path}")
    lines.append(f"  - Log: {summary.log_path}")
    if summary.diff_output_path:
        lines.append(f"  - Diff: {summary.diff_output_path} ({summary.diff_format})")
    lines.append("- Muscle normalization:")
    alias_label = summary.muscle_alias_file or "default aliases"
    acronym_label = ", ".join(sorted(summary.muscle_acronyms)) if summary.muscle_acronyms else "none"
    lines.append(f"  - Alias file: {alias_label}")
    lines.append(f"  - Acronyms: {acronym_label}")
    lines.append(
        f"- Mode: dryrun={'true' if summary.dryrun else 'false'}, fix_names={'true' if summary.fix_names else 'false'}"
    )
    lines.append("- Policies:")
    scalar_override_text = ", ".join(f"{field}={policy}" for field, policy in summary.scalar_overrides.items()) or "none"
    list_override_text = ", ".join(f"{field}={policy}" for field, policy in summary.list_overrides.items()) or "none"
    lines.append(f"  - Scalar policy: {summary.scalar_policy} (overrides: {scalar_override_text})")
    lines.append(f"  - List policy: {summary.list_policy} (overrides: {list_override_text})")
    lines.append("- Counts:")
    lines.append(f"  - Baseline DB rows: {summary.baseline_rows}")
    lines.append(f"  - CSV rows read: {summary.csv_rows}")
    lines.append(f"  - Consolidated records: {summary.consolidated_records}")
    lines.append(f"  - Inserted: {summary.inserted}")
    lines.append(f"  - Updated: {summary.updated}")
    lines.append(f"  - No-ops: {summary.no_ops}")
    lines.append(f"  - Conflicts: {summary.conflicts}")
    if summary.skipped_rows:
        lines.append(f"  - Skipped rows: {', '.join(summary.skipped_rows)}")
    lines.append(f"- Name-trim fix: {summary.name_trim_status}")
    if summary.name_trim_csv:
        lines.append(f"- Name-trim report: {summary.name_trim_csv}")
    lines.append(f"- Timing: {summary.duration_seconds:.2f} seconds")
    lines.append("")

    lines.append(f"### Top Updates (limit {summary.diff_top_k})")
    if summary.top_updates:
        for row in summary.top_updates:
            lines.append(f"- {row.exercise_name} | {summarize_row(row)}")
    else:
        lines.append("- None.")

    if summary.errors:
        lines.append("")
        lines.append("### Errors")
        lines.append("```")
        lines.append(summary.errors)
        lines.append("```")

    if summary.dryrun:
        lines.append("")
        lines.append("NO CHANGES APPLIED")

    with open(summary.log_path, "a", encoding="utf-8") as handle:
        if handle.tell() > 0:
            handle.write("\n\n")
        handle.write("\n".join(lines))


def parse_overrides(
    spec: Optional[str],
    allowed_fields: Iterable[str],
    allowed_policies: Set[str],
) -> Dict[str, str]:
    if not spec:
        return {}
    field_set = {field for field in allowed_fields}
    overrides: Dict[str, str] = {}
    for item in spec.split(","):
        token = item.strip()
        if not token:
            continue
        if "=" not in token:
            raise ValueError(f"Invalid override entry: {token}")
        field, policy = token.split("=", 1)
        field = field.strip()
        policy = policy.strip().lower()
        if field not in field_set:
            raise ValueError(f"Unknown field in override: {field}")
        if policy not in allowed_policies:
            raise ValueError(f"Unsupported policy '{policy}' for field {field}")
        overrides[field] = policy
    return overrides


def run_merge(
    csv_path: str,
    db_path: str,
    log_path: str,
    *,
    dryrun: bool,
    fix_names: bool,
    scalar_policy: str,
    scalar_overrides: Dict[str, str],
    list_policy: str,
    list_overrides: Dict[str, str],
    export_diff_path: Optional[str],
    diff_format: str,
    diff_top_k: int,
    alias_map: Optional[Dict[str, str]] = None,
    acronyms: Optional[Set[str]] = None,
    alias_file: Optional[str] = None,
) -> MergeSummary:
    summary = MergeSummary(
        timestamp=dt.datetime.now(),
        csv_path=os.path.abspath(csv_path),
        db_path=os.path.abspath(db_path),
        log_path=os.path.abspath(log_path),
        dryrun=dryrun,
        fix_names=fix_names,
        scalar_policy=scalar_policy,
        scalar_overrides=scalar_overrides,
        list_policy=list_policy,
        list_overrides=list_overrides,
        diff_top_k=diff_top_k,
    )
    summary.muscle_alias_file = os.path.abspath(alias_file) if alias_file else None
    start = time.time()
    policy_config = PolicyConfig(scalar_policy, scalar_overrides, list_policy, list_overrides)
    conn: Optional[sqlite3.Connection] = None

    muscle_alias_map = dict(alias_map) if alias_map is not None else load_muscle_alias_map(None)
    muscle_acronyms = set(acronyms) if acronyms is not None else set(MUSCLE_DEFAULT_ACRONYMS)
    summary.muscle_acronyms = set(muscle_acronyms)

    try:
        consolidated, skipped_rows, total_rows = load_csv_records(summary.csv_path)
        summary.csv_rows = total_rows
        summary.consolidated_records = len(consolidated)
        summary.skipped_rows = skipped_rows

        for record in consolidated.values():
            apply_muscle_normalization(record, muscle_alias_map, muscle_acronyms)

        if not os.path.exists(summary.db_path):
            raise FileNotFoundError(f"SQLite database not found: {summary.db_path}")

        if not summary.dryrun:
            summary.backup_path = backup_database(summary.db_path)

        conn = sqlite3.connect(summary.db_path)
        conn.row_factory = sqlite3.Row
        ensure_exercises_table(conn)

        summary.baseline_rows = conn.execute("SELECT COUNT(*) FROM exercises").fetchone()[0]
        conn.execute("BEGIN")

        for csv_record in consolidated.values():
            name = csv_record.get(PRIMARY_KEY)
            if not name:
                continue
            existing = fetch_existing_row(conn, name)
            if existing is None:
                record, row_diff = build_insert_record(csv_record, muscle_alias_map, muscle_acronyms)
                summary.diff_entries.append(row_diff)
                summary.inserted += 1
                if not summary.dryrun:
                    insert_row(conn, record)
            else:
                merged_record, row_diff, needs_update = merge_existing_record(
                    existing,
                    csv_record,
                    policy_config,
                    summary,
                    muscle_alias_map,
                    muscle_acronyms,
                )
                summary.diff_entries.append(row_diff)
                if row_diff.change_type == "Updated":
                    summary.updated += 1
                    if not summary.dryrun and needs_update:
                        update_row(conn, merged_record)
                else:
                    summary.no_ops += 1
        if summary.dryrun:
            conn.rollback()
        else:
            conn.commit()
    except Exception:
        if conn is not None:
            conn.rollback()
        summary.errors = traceback.format_exc()
    finally:
        if conn is not None:
            conn.close()

    summary.duration_seconds = time.time() - start

    if summary.diff_entries:
        updated_rows = [row for row in summary.diff_entries if row.change_type == "Updated"]
        sorted_updates = sorted(updated_rows, key=lambda row: len(row.columns_changed), reverse=True)
        summary.top_updates = sorted_updates[: diff_top_k if diff_top_k > 0 else 0]
    else:
        summary.top_updates = []

    if summary.errors is None and export_diff_path:
        try:
            export_diff(summary, export_diff_path, diff_format)
        except Exception:
            summary.errors = summary.errors or traceback.format_exc()

    if summary.fix_names:
        if summary.errors is None:
            status, csv_path = perform_name_trim(summary.db_path, summary.log_path, summary.dryrun)
            summary.name_trim_status = status
            summary.name_trim_csv = csv_path
        else:
            summary.name_trim_status = "skipped (error encountered)"

    append_log(summary)
    return summary


def run_selftest() -> None:
    import tempfile

    with tempfile.TemporaryDirectory() as tmpdir:
        db_path = os.path.join(tmpdir, "test.db")
        csv_path = os.path.join(tmpdir, "data.csv")
        log_path = os.path.join(tmpdir, "log.md")

        conn = sqlite3.connect(db_path)
        ensure_exercises_table(conn)
        conn.commit()
        conn.close()

        with open(csv_path, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(TARGET_COLUMNS)
            writer.writerow([
                "squat",
                "Quadriceps",
                "Glutes",
                "Hamstrings",
                "",
                "Basic",
                "",
                "",
                "",
                "Push",
                "Barbell",
                "Compound",
                "Beginner",
            ])
        summary = run_merge(
            csv_path,
            db_path,
            log_path,
            dryrun=False,
            fix_names=False,
            scalar_policy="db",
            scalar_overrides={},
            list_policy="union",
            list_overrides={},
            export_diff_path=None,
            diff_format="csv",
            diff_top_k=5,
        )
        assert summary.inserted == 1

        with open(csv_path, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(TARGET_COLUMNS)
            writer.writerow([
                "squat",
                "Quadriceps",
                "Glutes",
                "Hamstrings",
                "",
                "Basic",
                "Pronated",
                "",
                "",
                "Push",
                "Barbell",
                "Compound",
                "Beginner",
            ])
        summary = run_merge(
            csv_path,
            db_path,
            log_path,
            dryrun=False,
            fix_names=False,
            scalar_policy="db",
            scalar_overrides={},
            list_policy="union",
            list_overrides={},
            export_diff_path=None,
            diff_format="csv",
            diff_top_k=5,
        )
        assert summary.updated >= 1

        summary = run_merge(
            csv_path,
            db_path,
            log_path,
            dryrun=True,
            fix_names=True,
            scalar_policy="db",
            scalar_overrides={},
            list_policy="union",
            list_overrides={},
            export_diff_path=None,
            diff_format="csv",
            diff_top_k=5,
        )
        assert summary.dryrun
        assert summary.name_trim_status.startswith("name-trim") or summary.name_trim_status in {
            "no trim candidates",
            "name-trim fix skipped (dryrun)",
        }

    print("Self-test passed.")


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Merge exercise progress CSV into SQLite database.")
    parser.add_argument("--csv", default=DEFAULT_CSV_PATH, help="Path to the source CSV file")
    parser.add_argument("--db", default=DEFAULT_DB_PATH, help="Path to the target SQLite database")
    parser.add_argument("--log", default=DEFAULT_LOG_PATH, help="Path to the Markdown log file")
    parser.add_argument("--dryrun", action="store_true", help="Run without persisting changes")
    parser.add_argument("--fix_names", action="store_true", help="Trim exercise names post-merge when safe")
    parser.add_argument(
        "--scalar-policy",
        choices=sorted(ALLOWED_SCALAR_POLICIES),
        default="db",
        help="Default conflict policy for scalar fields",
    )
    parser.add_argument(
        "--scalar-overrides",
        help="Comma-separated per-field scalar policy overrides (e.g. equipment=csv)",
    )
    parser.add_argument(
        "--list-policy",
        choices=sorted(ALLOWED_LIST_POLICIES),
        default="union",
        help="Default policy for list fields",
    )
    parser.add_argument(
        "--list-overrides",
        help="Comma-separated per-field list policy overrides (e.g. grips=csv)",
    )
    parser.add_argument("--export-diff", help="Optional diff export file path")
    parser.add_argument(
        "--diff-format",
        choices=["csv", "json", "md", "xlsx"],
        default="csv",
        help="Format for diff export when --export-diff is provided",
    )
    parser.add_argument(
        "--diff-top-k",
        type=int,
        default=20,
        help="Top-K updated rows to emit to stdout and log",
    )
    parser.add_argument(
        "--alias-file",
        help="Optional YAML/JSON alias map for muscle name normalization",
    )
    parser.add_argument("--selftest", action="store_true", help="Run internal self-test suite and exit")
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> None:
    args = parse_args(argv)
    if args.selftest:
        run_selftest()
        return

    if args.diff_top_k < 0:
        print("--diff-top-k must be >= 0", file=sys.stderr)
        sys.exit(2)

    try:
        scalar_overrides = parse_overrides(args.scalar_overrides, SCALAR_COLUMNS, ALLOWED_SCALAR_POLICIES)
        list_overrides = parse_overrides(args.list_overrides, LIST_COLUMNS, ALLOWED_LIST_POLICIES)
    except ValueError as exc:
        print(f"Invalid override specification: {exc}", file=sys.stderr)
        sys.exit(2)

    alias_path = os.path.abspath(args.alias_file) if args.alias_file else None
    alias_map = load_muscle_alias_map(alias_path)
    acronyms = set(MUSCLE_DEFAULT_ACRONYMS)
    acronyms.update(get_muscle_alias_acronyms())

    summary = run_merge(
        args.csv,
        args.db,
        args.log,
        dryrun=args.dryrun,
        fix_names=args.fix_names,
        scalar_policy=args.scalar_policy,
        scalar_overrides=scalar_overrides,
        list_policy=args.list_policy,
        list_overrides=list_overrides,
        export_diff_path=args.export_diff,
        diff_format=args.diff_format,
        diff_top_k=args.diff_top_k,
        alias_map=alias_map,
        acronyms=acronyms,
        alias_file=alias_path,
    )

    if summary.errors:
        print("Merge failed; see log for details.", file=sys.stderr)
        print(summary.errors, file=sys.stderr)
        sys.exit(1)

    print(
        "Merge complete: "
        f"csv_rows={summary.csv_rows}, "
        f"consolidated={summary.consolidated_records}, "
        f"inserted={summary.inserted}, "
        f"updated={summary.updated}, "
        f"no_ops={summary.no_ops}, "
        f"conflicts={summary.conflicts}, "
        f"dryrun={summary.dryrun}"
    )
    if summary.backup_path:
        print(f"Database backup: {summary.backup_path}")
    if summary.diff_output_path:
        print(f"Diff exported to: {summary.diff_output_path} ({summary.diff_format})")
    alias_label = summary.muscle_alias_file or "default aliases"
    acronyms_label = ", ".join(sorted(summary.muscle_acronyms)) if summary.muscle_acronyms else "none"
    print(f"Muscle aliases: {alias_label}")
    print(f"Muscle acronyms: {acronyms_label}")
    if summary.diff_top_k > 0 and summary.top_updates:
        print(f"Top changes (limit {summary.diff_top_k}):")
        for index, row in enumerate(summary.top_updates, start=1):
            print(f"  {index}. {row.exercise_name} — {summarize_row(row)}")
    elif summary.diff_top_k > 0:
        print("No updated rows to report.")
    if summary.fix_names:
        print(f"Name-trim: {summary.name_trim_status}")
    print(f"Run log appended to: {summary.log_path}")


if __name__ == "__main__":
    main()
